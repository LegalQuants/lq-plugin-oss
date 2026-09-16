"""Calibration gate: downstream artifacts refuse without it, re-run on hash
drift, and a row's direction renders only when present."""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[4]
# The skill lives under a practice folder since the plugin split; find it by name.
SCRIPTS = next(ROOT.glob("skills/**/read-redline/scripts"))


def _load(name: str):
    # sibling imports (calibration_gate) resolve off the script dir, as they
    # do when the scripts run as `python3 scripts/<name>.py`
    if str(SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


def _write_run(tmp_path: Path, *, with_direction: bool = True) -> dict[str, Path]:
    extract = {
        "source_pdf": "/tmp/nda.pdf",
        "pages": [{"number": 0, "width": 612, "height": 792, "rotation": 0}],
        "calibration": {"roles": {"16711680": "del", "255": "ins"}},
        "pairs": [
            {
                "pair_id": "p-0001",
                "page": 0,
                "old_text": "thirty days",
                "new_text": "ten days",
                "changed": True,
            }
        ],
        "move_annotations": [],
        "quarantined": [],
        "stats": {"pages": 1},
    }
    row = {
        "provision": "Notice period",
        "source_pair_ids": ["p-0001"],
        "comment": "Notice shrinks from thirty days to ten.",
        "materiality": "Medium",
    }
    if with_direction:
        row["direction"] = "favours-them"
    themes = {
        "themes": [
            {"label": "Notice", "summary": "Tightening notice.", "row_refs": [0]}
        ],
        "housekeeping": {"count": 0, "row_refs": []},
        "unclustered": [],
        "coverage": {
            "rows_total": 1,
            "rows_themed": 1,
            "rows_housekeeping": 0,
            "rows_unclustered": 0,
        },
    }
    paths = {}
    for name, data in (("extract", extract), ("rows", [row]), ("themes", themes)):
        path = tmp_path / f"{name}.json"
        path.write_text(json.dumps(data), encoding="utf-8")
        paths[name] = path
    return paths


def _issues_list(tmp_path: Path, paths: dict[str, Path]) -> Path:
    mil = _load("make_issues_list")
    out = tmp_path / "Issues List.docx"
    mil.main(str(paths["extract"]), str(paths["rows"]), str(paths["themes"]), str(out))
    return out


def _provision_cell(out: Path) -> str:
    from docx import Document

    table = Document(out).tables[0]
    return table.rows[2].cells[1].text  # header, theme band, then the issue row


def test_issues_list_refuses_without_gate(tmp_path: Path) -> None:
    pytest.importorskip("docx")  # make_issues_list.py exits at import without it
    mil = _load("make_issues_list")
    paths = _write_run(tmp_path)
    with pytest.raises(SystemExit) as exc:
        mil.main(
            str(paths["extract"]),
            str(paths["rows"]),
            str(paths["themes"]),
            str(tmp_path / "out.docx"),
        )
    assert exc.value.code != 0
    assert "calibration.confirmed.json" in str(exc.value.code)


def test_issues_list_runs_with_declared_default(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    gate = _load("calibration_gate")
    paths = _write_run(tmp_path)
    gate.write_gate(
        paths["extract"], "declared-default", "batch run; no reviewer present"
    )
    out = _issues_list(tmp_path, paths)
    from docx import Document

    receipt = "\n".join(p.text for p in Document(out).paragraphs)
    assert "declared default" in receipt
    assert "batch run; no reviewer present" in receipt


def test_issues_list_refuses_on_hash_drift(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    gate = _load("calibration_gate")
    mil = _load("make_issues_list")
    paths = _write_run(tmp_path)
    gate.write_gate(paths["extract"], "confirmed")
    extract = json.loads(paths["extract"].read_text(encoding="utf-8"))
    extract["quarantined"].append({"page": 0, "reason": "visible-mark-not-extracted"})
    paths["extract"].write_text(json.dumps(extract), encoding="utf-8")
    with pytest.raises(SystemExit) as exc:
        mil.main(
            str(paths["extract"]),
            str(paths["rows"]),
            str(paths["themes"]),
            str(tmp_path / "out.docx"),
        )
    assert exc.value.code != 0
    assert "hash" in str(exc.value.code)


def test_direction_renders_beside_tier_with_location(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    gate = _load("calibration_gate")
    paths = _write_run(tmp_path, with_direction=True)
    gate.write_gate(paths["extract"], "confirmed")
    cell = _provision_cell(_issues_list(tmp_path, paths))
    assert "favours-them" in cell
    assert "Medium" in cell
    assert "p. 1" in cell  # parser pages are 0-indexed; display is 1-based


def test_direction_absent_renders_cleanly(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    gate = _load("calibration_gate")
    paths = _write_run(tmp_path, with_direction=False)
    gate.write_gate(paths["extract"], "confirmed")
    cell = _provision_cell(_issues_list(tmp_path, paths))
    assert "Medium" in cell
    assert "·" not in cell
    assert "favours" not in cell


def test_declared_default_requires_a_reason(tmp_path: Path) -> None:
    gate = _load("calibration_gate")
    paths = _write_run(tmp_path)
    with pytest.raises(SystemExit) as exc:
        gate.write_gate(paths["extract"], "declared-default")
    assert exc.value.code != 0
    assert "reason" in str(exc.value.code)


def test_docx_comment_head_prefixes_direction() -> None:
    ann = _load("annotate_docx")
    row = {
        "provision": "Liability cap",
        "materiality": "High",
        "comment": "The cap halves.",
    }
    assert ann.build_comment(row) == "[High] Liability cap — The cap halves."
    assert ann.build_comment(row | {"direction": "favours-us"}) == (
        "[favours-us · High] Liability cap — The cap halves."
    )


def test_pdf_subject_prefixes_direction() -> None:
    pytest.importorskip("pypdf")  # annotate_pdf.py exits at import without it
    ann = _load("annotate_pdf")
    by_id = {"p-0001": {"page": 0, "boxes": {"provision": [[0, 0, 10, 10]]}}}
    heights = {0: 792.0}
    row = {
        "provision": "Liability cap",
        "materiality": "High",
        "comment": "The cap halves.",
        "source_pair_ids": ["p-0001"],
    }
    plain, _ = ann.collect_annotations([row], by_id, heights)
    assert plain[0][0]["subject"] == "High — Liability cap"
    directed, _ = ann.collect_annotations(
        [row | {"direction": "favours-us"}], by_id, heights
    )
    assert directed[0][0]["subject"] == "favours-us · High — Liability cap"
